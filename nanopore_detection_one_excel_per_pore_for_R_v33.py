"""
Nanopore event detection v33: one Excel workbook per pore, prepared for R.

Expected input layout
---------------------
DATA_ROOT/
    Data_001_..._aHL_tauThr217/
        recording.abf
        IV Curve/             <- never scanned
        other channels/       <- never scanned
        baseline/             <- never scanned
    Data_002_..._aHL_AT8/
        recording.abf

Every immediate child directory of DATA_ROOT that contains at least one ABF
directly inside it is treated as one physical pore. The script never searches
recursively. Pores are grouped into peptide output folders, but every physical
pore receives its own Excel workbook directly inside the peptide folder.
Python creates no plots, CSV files, or per-pore output directories.

The event detector and output filters preserve the settings from v31. A
deterministically selected 2 s current segment is also exported for every pore
as a trace_Pn worksheet for later plotting in R.
"""

from __future__ import annotations

import gc
import json
import os
import re
import tempfile
import traceback
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import pyabf
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from scipy.signal import medfilt
from scipy.stats import median_abs_deviation


# ---------------------------------------------------------------------------
# USER SETTINGS
# ---------------------------------------------------------------------------

# Point this to the Recordings directory shown in Windows Explorer.
DATA_ROOT = r"C:\Users\mbech\Documents\Nouveaux pores à la date du 1 sept"

OUTPUT_ROOT_NAME = "RobustDetection_R_Excel_per_pore"
CHANNEL = 0
DETECTION_CHUNK_DURATION_S = 5.0

# Canonical peptide name: aliases that may occur in a pore-folder or ABF name.
# Matching ignores spaces, underscores, hyphens, and letter case.
PEPTIDE_ALIASES: Dict[str, Tuple[str, ...]] = {
    "AT8": (
        "AT8",
        "tauAT8",
        "taupSer202pThr205",
        "taupSer202_pThr205"
        "taupSer202_Thr205",
        "taupSer202Thr205",
        "pSer202_Thr205",
        "pSer202Thr205",
        "Tau pSer202 Thr205",
    ),
    "taupThr217": (
        "taupThr217",
        "taup217",
        "Tau pThr217",
        "pThr217",
    ),
    "tauThr217": (
        "tauThr217",
        "Tau Thr217",
        "Thr217",
    ),
    "taupThr181": (
        "taupThr181",
        "taup181",
        "Tau pThr181",
        "pThr181",
    ),
    "tauThr181": (
        "tauThr181",
        "Tau Thr181",
        "Thr181",
    ),
    "taupSer262": (
        "taupSer262",
        "taupS262",
        "Tau pSer262",
        "pSer262",
    ),
    "tauSer262": (
        "tauSer262",
        "Tau Ser262",
        "Ser262",
    ),
    "tauSer202_Thr205": (
        "tauSer202_Thr205",
        "tauSer202Thr205",
        "Ser202_Thr205",
        "Ser202Thr205",
        "Tau Ser202 Thr205",
    ),
    "asynpS87": (
        "asynpS87",
        "alphaSynpS87",
        "pS87",
    ),
    "asynY125": (
        "asynY125",
        "alphaSynY125",
        "Y125",
    ),
    "asynpY125": (
        "asynpY125",
        "alphaSynpY125",
        "pY125",
    ),
    "asynY39": (
        "asynY39",
        "alphaSynY39",
        "Y39",
    ),
}

# Use this only when a folder name cannot be classified from PEPTIDE_ALIASES.
# The key must be the exact immediate folder name inside DATA_ROOT.
FOLDER_TO_PEPTIDE_OVERRIDES: Dict[str, str] = {
    # "Data_999_example_folder": "AT8",
}


@dataclass
class DetectConfig:
    # Analysed part of every ABF sweep
    restrict_time_window: bool = True
    window_start_s: float = 0.0
    window_duration_s: float = 300.0

    # Local baseline and noise
    baseline_win_s: float = 0.50
    sigma_win_s: float = 0.25
    sigma_update_step_samples: int = 1000

    # Hysteresis event detection
    sigma_enter: float = 3.5
    sigma_exit: float = 1.5
    min_dwell_ms: float = 0.05
    max_dwell_ms: float = 500.0
    merge_gap_ms: float = 0.10
    min_core_samples: int = 2
    min_exit_samples: int = 2

    # Per-sweep QC
    enable_between_sweep_qc: bool = True
    qc_group_columns: Tuple[str, ...] = ("analyte",)
    qc_mad_z_threshold: float = 3.5
    qc_min_sweeps_per_group: int = 3
    qc_min_mad_floor: float = 1e-12
    qc_metrics: Tuple[str, ...] = (
        "event_rate_Hz",
        "median_dwell_total_ms",
        "median_I_over_I0_pct",
        "median_blockade_depth_mean_pA",
        "baseline_median_pA",
        "sigma_median_pA",
    )
    flat_trace_std_pA: float = 1e-3
    large_baseline_drift_pA: float = 10.0
    high_event_fraction: float = 0.30

    # Events exported in the main events worksheet
    filter_events_before_output: bool = True
    min_dwell_filter_ms: float = 0.2
    max_dwell_filter_ms: float = 15.0
    min_depth_filter_pA: float = 0.0
    max_depth_filter_pA: float = np.inf
    depth_filter_column: str = "blockade_depth_mean_pA"
    filter_by_i_over_i0_pct: bool = False
    min_i_over_i0_filter_pct: float = 0.0
    max_i_over_i0_filter_pct: float = 40.0

    # Representative current segment exported for R
    representative_trace_duration_s: float = 2.0
    representative_trace_step_s: float = 0.5
    representative_trace_min_events: int = 3
    representative_trace_max_event_fraction: float = 0.20


CFG = DetectConfig()


BASE_EVENT_COLUMNS = [
    "event_id",
    "file_name",
    "pore_id",
    "analyte",
    "channel",
    "sweep",
    "global_start_idx",
    "global_end_idx",
    "start_idx_in_sweep",
    "end_idx_in_sweep",
    "start_s_in_sweep",
    "end_s_in_sweep",
    "start_s_global",
    "end_s_global",
    "dwell_core_ms",
    "dwell_total_ms",
    "baseline_local_pA",
    "sigma_local_pA",
    "threshold_enter_pA",
    "threshold_exit_pA",
    "I_mean_pA",
    "I_median_pA",
    "I_min_pA",
    "blockade_depth_mean_pA",
    "blockade_depth_median_pA",
    "blockade_depth_min_pA",
    "I_over_I0_pct",
    "iteration_detected",
]

EVENT_EXTRA_COLUMNS = [
    "raw_event_id",
    "experiment_index",
    "source_pore_folder",
]

FINAL_EVENT_COLUMNS = BASE_EVENT_COLUMNS + EVENT_EXTRA_COLUMNS
EXCEL_MAX_DATA_ROWS = 1_048_575


# ---------------------------------------------------------------------------
# DISCOVERY: ONLY IMMEDIATE PORE FOLDERS AND DIRECT ABFs
# ---------------------------------------------------------------------------

def natural_sort_key(value: object) -> Tuple[object, ...]:
    return tuple(
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", str(value))
    )


def safe_name(value: object) -> str:
    text = re.sub(r'[<>:"/\\|?*]+', "_", str(value)).strip(" .")
    return text or "unnamed"


def normalize_match_text(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def infer_analyte(
    folder_name: str,
    direct_abfs: Sequence[Path],
) -> Tuple[Optional[str], str]:
    if folder_name in FOLDER_TO_PEPTIDE_OVERRIDES:
        analyte = FOLDER_TO_PEPTIDE_OVERRIDES[folder_name]
        if analyte not in PEPTIDE_ALIASES:
            return None, f"override points to unknown peptide: {analyte}"
        return analyte, "manual folder override"

    combined = " ".join([folder_name] + [path.stem for path in direct_abfs])
    normalized = normalize_match_text(combined)

    matches: List[Tuple[int, str, str]] = []
    for analyte, aliases in PEPTIDE_ALIASES.items():
        for alias in aliases:
            normalized_alias = normalize_match_text(alias)
            if normalized_alias and normalized_alias in normalized:
                matches.append((len(normalized_alias), analyte, alias))

    if not matches:
        return None, "no peptide alias found in folder or direct ABF name"

    matches.sort(key=lambda item: (-item[0], item[1], item[2]))
    best_length = matches[0][0]
    best_analytes = sorted({item[1] for item in matches if item[0] == best_length})
    if len(best_analytes) > 1:
        return None, "ambiguous longest aliases: " + ", ".join(best_analytes)

    best = matches[0]
    return best[1], f"matched alias: {best[2]}"


def discover_recordings(
    data_root: Path,
) -> Tuple[Dict[str, List[Dict[str, object]]], pd.DataFrame]:
    experiments_by_analyte: Dict[str, List[Dict[str, object]]] = {}
    audit_rows: List[Dict[str, object]] = []

    folders = sorted(
        (
            path
            for path in data_root.iterdir()
            if path.is_dir() and path.name != OUTPUT_ROOT_NAME
        ),
        key=lambda path: natural_sort_key(path.name),
    )

    pending: Dict[str, List[Tuple[Path, List[Path], str]]] = {}

    for folder in folders:
        direct_abfs = sorted(
            (
                path
                for path in folder.iterdir()
                if path.is_file() and path.suffix.lower() == ".abf"
            ),
            key=lambda path: natural_sort_key(path.name),
        )

        if not direct_abfs:
            audit_rows.append({
                "source_pore_folder": folder.name,
                "status": "skipped",
                "reason": "no ABF directly inside folder; subfolders were not scanned",
                "analyte": "",
                "pore_id": "",
                "abf_file": "",
                "abf_path": "",
            })
            continue

        analyte, reason = infer_analyte(folder.name, direct_abfs)
        if analyte is None:
            for abf_path in direct_abfs:
                audit_rows.append({
                    "source_pore_folder": folder.name,
                    "status": "skipped",
                    "reason": reason,
                    "analyte": "",
                    "pore_id": "",
                    "abf_file": abf_path.name,
                    "abf_path": str(abf_path),
                })
            continue

        pending.setdefault(analyte, []).append((folder, direct_abfs, reason))

    experiment_index = 0
    for analyte in sorted(pending, key=natural_sort_key):
        pore_folders = sorted(pending[analyte], key=lambda item: natural_sort_key(item[0].name))
        experiments: List[Dict[str, object]] = []

        for pore_number, (folder, direct_abfs, reason) in enumerate(pore_folders, start=1):
            pore_id = f"P{pore_number}"
            for abf_path in direct_abfs:
                experiment_index += 1
                experiment = {
                    "experiment_index": experiment_index,
                    "analyte": analyte,
                    "pore_id": pore_id,
                    "source_pore_folder": folder.name,
                    "abf_path": str(abf_path),
                    "channel": CHANNEL,
                }
                experiments.append(experiment)
                audit_rows.append({
                    "source_pore_folder": folder.name,
                    "status": "used",
                    "reason": reason,
                    "analyte": analyte,
                    "pore_id": pore_id,
                    "abf_file": abf_path.name,
                    "abf_path": str(abf_path),
                })

        experiments_by_analyte[analyte] = experiments

    audit = pd.DataFrame(audit_rows, columns=[
        "source_pore_folder",
        "status",
        "reason",
        "analyte",
        "pore_id",
        "abf_file",
        "abf_path",
    ])
    return experiments_by_analyte, audit


# ---------------------------------------------------------------------------
# EVENT DETECTION (SAME CORE LOGIC AND SETTINGS AS v31)
# ---------------------------------------------------------------------------

def empty_events_df() -> pd.DataFrame:
    return pd.DataFrame(columns=BASE_EVENT_COLUMNS)


def safe_median(values: np.ndarray) -> float:
    return float(np.median(values)) if values.size else np.nan


def safe_mean(values: np.ndarray) -> float:
    return float(np.mean(values, dtype=np.float32)) if values.size else np.nan


def safe_std(values: np.ndarray) -> float:
    return float(np.std(values, dtype=np.float32)) if values.size else np.nan


def contiguous_true_regions(mask: np.ndarray) -> List[Tuple[int, int]]:
    if mask.size == 0:
        return []

    changes = np.diff(mask.astype(np.int8))
    starts = list(np.where(changes == 1)[0] + 1)
    ends = list(np.where(changes == -1)[0])

    if mask[0]:
        starts.insert(0, 0)
    if mask[-1]:
        ends.append(len(mask) - 1)

    return list(zip(starts, ends))


def merge_close_regions(
    regions: List[Tuple[int, int]],
    max_gap_points: int,
) -> List[Tuple[int, int]]:
    if not regions:
        return []

    merged = [list(regions[0])]
    for start, end in regions[1:]:
        gap = start - merged[-1][1] - 1
        if gap <= max_gap_points:
            merged[-1][1] = end
        else:
            merged.append([start, end])

    return [(int(start), int(end)) for start, end in merged]


def estimate_open_pore_stats(
    values: np.ndarray,
    nbins: int = 400,
    max_points: int = 50_000,
) -> Tuple[float, float]:
    values = np.asarray(values, dtype=np.float32)
    values = values[np.isfinite(values)]

    if values.size == 0:
        return np.nan, np.nan
    if values.size < 10:
        return safe_mean(values), safe_std(values) + 1e-12

    if values.size > max_points:
        indices = np.linspace(0, values.size - 1, max_points, dtype=np.int64)
        work = values[indices]
    else:
        work = values

    histogram, edges = np.histogram(work, bins=nbins)
    peak_index = int(np.argmax(histogram))
    center = np.float32(0.5 * (edges[peak_index] + edges[peak_index + 1]))
    bin_width = np.float32(edges[1] - edges[0])
    select = (work >= center - 4.0 * bin_width) & (work <= center + 4.0 * bin_width)
    peak = work[select] if np.any(select) else work

    i0 = float(np.mean(peak, dtype=np.float32))
    sigma0 = float(np.std(peak, dtype=np.float32)) + 1e-12
    return i0, sigma0


def sample_abf_for_i0(
    abf: pyabf.ABF,
    channel: int,
    max_points_total: int = 50_000,
) -> Tuple[float, float, float]:
    if channel < 0 or channel >= abf.channelCount:
        raise ValueError(
            f"Invalid channel {channel}; ABF contains {abf.channelCount} channel(s)."
        )

    samples: List[np.ndarray] = []
    points_per_sweep = max(1, int(np.ceil(max_points_total / max(1, abf.sweepCount))))

    for sweep in range(abf.sweepCount):
        abf.setSweep(sweep, channel=channel)
        values = np.asarray(abf.sweepY, dtype=np.float32)
        if not values.size:
            continue
        take = min(points_per_sweep, len(values))
        indices = np.linspace(0, len(values) - 1, take, dtype=np.int64)
        samples.append(values[indices])

    if not samples:
        return np.nan, np.nan, 1.0

    sample = np.concatenate(samples).astype(np.float32, copy=False)
    median_absolute_current = float(np.nanmedian(np.abs(sample)))
    unit_scale = 1000.0 if 0.01 <= median_absolute_current < 5.0 else 1.0
    if unit_scale != 1.0:
        sample = sample * np.float32(unit_scale)
        print(
            f"  Unit check: signal looked like nA "
            f"(median |I|={median_absolute_current:.4g}); converted to pA."
        )
    else:
        print(
            f"  Unit check: signal kept as pA "
            f"(median |I|={median_absolute_current:.4g})."
        )

    i0, sigma0 = estimate_open_pore_stats(sample)
    return i0, sigma0, unit_scale


def rolling_median_baseline(
    values: np.ndarray,
    fs_hz: float,
    window_s: float,
) -> np.ndarray:
    values = np.asarray(values, dtype=np.float32)
    width = int(round(window_s * fs_hz))

    if width < 3:
        return np.full(values.shape, np.median(values), dtype=np.float32)
    if width % 2 == 0:
        width += 1
    if width >= len(values):
        return np.full(values.shape, np.median(values), dtype=np.float32)

    return medfilt(values, kernel_size=width).astype(np.float32, copy=False)


def rolling_mad_sigma(
    values: np.ndarray,
    baseline: np.ndarray,
    fs_hz: float,
    window_s: float,
    step_samples: int = 1000,
) -> np.ndarray:
    values = np.asarray(values, dtype=np.float32)
    baseline = np.asarray(baseline, dtype=np.float32)
    n_points = len(values)

    if n_points == 0:
        return np.array([], dtype=np.float32)

    step_samples = max(1, int(step_samples))
    width = int(round(window_s * fs_hz))

    if width < 3:
        sigma = float(np.std(values - baseline) + 1e-12)
        return np.full(values.shape, sigma, dtype=np.float32)
    if width % 2 == 0:
        width += 1

    residual = np.abs(values - baseline).astype(np.float32, copy=False)
    if width >= n_points:
        sigma = 1.4826 * float(np.median(residual)) + 1e-12
        return np.full(values.shape, sigma, dtype=np.float32)

    half_width = width // 2
    centers = np.arange(0, n_points, step_samples, dtype=np.int64)
    if centers.size == 0 or centers[0] != 0:
        centers = np.insert(centers, 0, 0)
    if centers[-1] != n_points - 1:
        centers = np.append(centers, n_points - 1)

    sparse_sigma = np.empty(len(centers), dtype=np.float32)
    for index, center in enumerate(centers):
        low = max(0, int(center) - half_width)
        high = min(n_points, int(center) + half_width + 1)
        sparse_sigma[index] = np.float32(
            1.4826 * float(np.median(residual[low:high])) + 1e-12
        )

    sigma = np.empty(n_points, dtype=np.float32)
    for index in range(len(centers) - 1):
        center_0 = int(centers[index])
        center_1 = int(centers[index + 1])
        if center_1 <= center_0:
            continue
        sigma[center_0:center_1] = np.linspace(
            float(sparse_sigma[index]),
            float(sparse_sigma[index + 1]),
            center_1 - center_0,
            endpoint=False,
            dtype=np.float32,
        )

    sigma[int(centers[-1]):] = sparse_sigma[-1]
    sigma[sigma < 1e-12] = np.float32(1e-12)
    return sigma


def derive_sample_params(
    fs_hz: float,
    cfg: DetectConfig,
) -> Dict[str, int]:
    return {
        "baseline_win_pts": max(1, int(round(cfg.baseline_win_s * fs_hz))),
        "sigma_win_pts": max(1, int(round(cfg.sigma_win_s * fs_hz))),
        "min_dwell_pts": max(
            1, int(round((cfg.min_dwell_ms / 1000.0) * fs_hz))
        ),
        "max_dwell_pts": max(
            1, int(round((cfg.max_dwell_ms / 1000.0) * fs_hz))
        ),
        "merge_gap_pts": max(
            0, int(round((cfg.merge_gap_ms / 1000.0) * fs_hz))
        ),
    }


def restrict_sweep_time_window(
    values: np.ndarray,
    fs_hz: float,
    cfg: DetectConfig,
) -> Tuple[np.ndarray, int, int]:
    if not cfg.restrict_time_window:
        return values, 0, len(values)

    start = max(0, int(round(cfg.window_start_s * fs_hz)))
    start = min(start, len(values))
    if start >= len(values):
        return values[0:0], start, start

    end = int(round((cfg.window_start_s + cfg.window_duration_s) * fs_hz))
    end = max(start + 1, min(end, len(values)))
    return values[start:end], start, end


def filter_events_for_outputs(
    events: pd.DataFrame,
    cfg: DetectConfig,
) -> pd.DataFrame:
    if events.empty or not cfg.filter_events_before_output:
        return events.copy()

    if cfg.depth_filter_column not in events.columns:
        raise ValueError(
            f"depth_filter_column={cfg.depth_filter_column!r} is missing."
        )

    mask = (
        (events["dwell_total_ms"] >= cfg.min_dwell_filter_ms)
        & (events["dwell_total_ms"] <= cfg.max_dwell_filter_ms)
        & (events[cfg.depth_filter_column] >= cfg.min_depth_filter_pA)
        & (events[cfg.depth_filter_column] <= cfg.max_depth_filter_pA)
    )

    if cfg.filter_by_i_over_i0_pct:
        mask &= (
            (events["I_over_I0_pct"] >= cfg.min_i_over_i0_filter_pct)
            & (events["I_over_I0_pct"] <= cfg.max_i_over_i0_filter_pct)
        )

    return events.loc[mask].reset_index(drop=True)


def hysteresis_detect_regions(
    values: np.ndarray,
    baseline: np.ndarray,
    sigma: np.ndarray,
    fs_hz: float,
    cfg: DetectConfig,
) -> Tuple[List[Tuple[int, int]], np.ndarray, np.ndarray]:
    threshold_enter = baseline - cfg.sigma_enter * sigma
    threshold_exit = baseline - cfg.sigma_exit * sigma

    below_enter = values < threshold_enter
    below_exit = values < threshold_exit
    regions: List[Tuple[int, int]] = []

    in_event = False
    start: Optional[int] = None
    core_count = 0
    exit_count = 0

    for index in range(len(values)):
        if not in_event:
            if below_enter[index]:
                if start is None:
                    start = index
                    core_count = 1
                else:
                    core_count += 1

                if core_count >= cfg.min_core_samples:
                    in_event = True
                    start = index if start is None else start
                    exit_count = 0
            else:
                start = None
                core_count = 0
        else:
            if below_exit[index]:
                exit_count = 0
            else:
                exit_count += 1
                if exit_count >= cfg.min_exit_samples:
                    end = index - cfg.min_exit_samples
                    event_start = index if start is None else start
                    regions.append((event_start, max(event_start, end)))
                    in_event = False
                    start = None
                    core_count = 0
                    exit_count = 0

    if in_event and start is not None:
        regions.append((start, len(values) - 1))

    params = derive_sample_params(fs_hz, cfg)
    regions = merge_close_regions(regions, params["merge_gap_pts"])
    return regions, threshold_enter, threshold_exit


def measure_events(
    values: np.ndarray,
    baseline: np.ndarray,
    sigma: np.ndarray,
    threshold_enter: np.ndarray,
    threshold_exit: np.ndarray,
    regions: List[Tuple[int, int]],
    fs_hz: float,
    cfg: DetectConfig,
    file_name: str,
    analyte: str,
    channel: int,
    sweep: int,
    global_offset: int,
    sweep_window_start_idx: int,
    i0_global: float,
) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    params = derive_sample_params(fs_hz, cfg)

    for start, end in regions:
        segment = values[start:end + 1]
        baseline_segment = baseline[start:end + 1]
        sigma_segment = sigma[start:end + 1]
        enter_segment = threshold_enter[start:end + 1]
        exit_segment = threshold_exit[start:end + 1]

        dwell_points = end - start + 1
        if dwell_points < params["min_dwell_pts"]:
            continue
        if dwell_points > params["max_dwell_pts"]:
            continue

        dwell_total_ms = dwell_points / fs_hz * 1000.0
        core_regions = contiguous_true_regions(segment < enter_segment)
        if core_regions:
            core_start = start + core_regions[0][0]
            core_end = start + core_regions[-1][1]
            dwell_core_ms = (core_end - core_start + 1) / fs_hz * 1000.0
        else:
            dwell_core_ms = dwell_total_ms

        start_in_sweep = sweep_window_start_idx + start
        end_in_sweep = sweep_window_start_idx + end
        baseline_local = safe_median(baseline_segment)
        sigma_local = safe_median(sigma_segment)
        current_mean = safe_mean(segment)
        current_median = safe_median(segment)
        current_min = float(np.min(segment))

        rows.append({
            "event_id": np.nan,
            "file_name": file_name,
            "pore_id": np.nan,
            "analyte": analyte,
            "channel": channel,
            "sweep": sweep,
            "global_start_idx": int(global_offset + start_in_sweep),
            "global_end_idx": int(global_offset + end_in_sweep),
            "start_idx_in_sweep": int(start_in_sweep),
            "end_idx_in_sweep": int(end_in_sweep),
            "start_s_in_sweep": float(start_in_sweep / fs_hz),
            "end_s_in_sweep": float(end_in_sweep / fs_hz),
            "start_s_global": float((global_offset + start_in_sweep) / fs_hz),
            "end_s_global": float((global_offset + end_in_sweep) / fs_hz),
            "dwell_core_ms": float(dwell_core_ms),
            "dwell_total_ms": float(dwell_total_ms),
            "baseline_local_pA": baseline_local,
            "sigma_local_pA": sigma_local,
            "threshold_enter_pA": safe_median(enter_segment),
            "threshold_exit_pA": safe_median(exit_segment),
            "I_mean_pA": current_mean,
            "I_median_pA": current_median,
            "I_min_pA": current_min,
            "blockade_depth_mean_pA": baseline_local - current_mean,
            "blockade_depth_median_pA": baseline_local - current_median,
            "blockade_depth_min_pA": baseline_local - current_min,
            "I_over_I0_pct": (
                float(100.0 * current_mean / i0_global)
                if np.isfinite(i0_global) and i0_global != 0
                else np.nan
            ),
            "iteration_detected": 1,
        })

    if not rows:
        return empty_events_df()

    return pd.DataFrame(rows)[BASE_EVENT_COLUMNS]


def detect_chunk(
    values: np.ndarray,
    fs_hz: float,
    cfg: DetectConfig,
    file_name: str,
    analyte: str,
    channel: int,
    sweep: int,
    global_offset: int,
    sweep_window_start_idx: int,
    i0_global: float,
) -> Tuple[
    pd.DataFrame,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    Dict[str, bool],
]:
    flags = {
        "flag_flat_trace": False,
        "flag_large_baseline_drift": False,
        "flag_high_event_fraction": False,
        "flag_no_events": False,
    }

    if not len(values):
        flags["flag_no_events"] = True
        empty = np.array([], dtype=np.float32)
        return empty_events_df(), empty, empty, empty, flags

    if np.std(values) < cfg.flat_trace_std_pA:
        flags["flag_flat_trace"] = True

    baseline = rolling_median_baseline(values, fs_hz, cfg.baseline_win_s)
    sigma = rolling_mad_sigma(
        values,
        baseline,
        fs_hz,
        cfg.sigma_win_s,
        cfg.sigma_update_step_samples,
    )
    regions, threshold_enter, threshold_exit = hysteresis_detect_regions(
        values,
        baseline,
        sigma,
        fs_hz,
        cfg,
    )
    events = measure_events(
        values=values,
        baseline=baseline,
        sigma=sigma,
        threshold_enter=threshold_enter,
        threshold_exit=threshold_exit,
        regions=regions,
        fs_hz=fs_hz,
        cfg=cfg,
        file_name=file_name,
        analyte=analyte,
        channel=channel,
        sweep=sweep,
        global_offset=global_offset,
        sweep_window_start_idx=sweep_window_start_idx,
        i0_global=i0_global,
    )

    event_fraction = 0.0
    if not events.empty:
        event_samples = (
            events["end_idx_in_sweep"]
            - events["start_idx_in_sweep"]
            + 1
        ).sum()
        event_fraction = float(event_samples / len(values))

    flags["flag_high_event_fraction"] = event_fraction > cfg.high_event_fraction
    flags["flag_large_baseline_drift"] = (
        float(np.max(baseline) - np.min(baseline))
        > cfg.large_baseline_drift_pA
    )
    flags["flag_no_events"] = events.empty
    return events, baseline, sigma, threshold_enter, flags


# ---------------------------------------------------------------------------
# REPRESENTATIVE TRACE SELECTION
# ---------------------------------------------------------------------------

TRACE_TIER_LABELS = {
    0: "at least target filtered events and acceptable event fraction",
    1: "at least one filtered event and acceptable event fraction",
    2: "no filtered event but acceptable event fraction",
    3: "fallback: event fraction above requested maximum",
}


def candidate_is_better(
    candidate: Dict[str, object],
    current: Optional[Dict[str, object]],
) -> bool:
    if current is None:
        return True
    return tuple(candidate["selection_rank"]) < tuple(current["selection_rank"])


def update_best_trace_from_chunk(
    current_best: Optional[Dict[str, object]],
    values: np.ndarray,
    baseline: np.ndarray,
    sigma: np.ndarray,
    threshold_enter: np.ndarray,
    events: pd.DataFrame,
    fs_hz: float,
    cfg: DetectConfig,
    analyte: str,
    pore_id: str,
    source_pore_folder: str,
    experiment_index: int,
    file_name: str,
    sweep: int,
    chunk_start_idx_in_sweep: int,
) -> Optional[Dict[str, object]]:
    if not len(values):
        return current_best

    target_points = max(1, int(round(cfg.representative_trace_duration_s * fs_hz)))
    step_points = max(1, int(round(cfg.representative_trace_step_s * fs_hz)))

    if len(values) <= target_points:
        starts = [0]
    else:
        starts = list(range(0, len(values) - target_points + 1, step_points))
        final_start = len(values) - target_points
        if starts[-1] != final_start:
            starts.append(final_start)

    filtered_events = filter_events_for_outputs(events, cfg)

    for local_start in starts:
        local_end = min(len(values), local_start + target_points)
        absolute_start = chunk_start_idx_in_sweep + local_start
        absolute_end = chunk_start_idx_in_sweep + local_end

        if filtered_events.empty:
            selected_events = filtered_events
        else:
            event_centers = 0.5 * (
                filtered_events["start_idx_in_sweep"].to_numpy(dtype=float)
                + filtered_events["end_idx_in_sweep"].to_numpy(dtype=float)
            )
            selected_events = filtered_events.loc[
                (event_centers >= absolute_start)
                & (event_centers < absolute_end)
            ]

        n_events = int(len(selected_events))
        event_samples = 0
        for row in selected_events.itertuples(index=False):
            overlap_start = max(absolute_start, int(row.start_idx_in_sweep))
            overlap_end = min(absolute_end - 1, int(row.end_idx_in_sweep))
            if overlap_end >= overlap_start:
                event_samples += overlap_end - overlap_start + 1

        window_points = max(1, local_end - local_start)
        event_fraction = float(event_samples / window_points)
        current_window = values[local_start:local_end]
        baseline_window = baseline[local_start:local_end]
        sigma_window = sigma[local_start:local_end]
        threshold_window = threshold_enter[local_start:local_end]

        baseline_drift = float(
            np.percentile(baseline_window, 95)
            - np.percentile(baseline_window, 5)
        )
        noise_mad = float(
            1.4826 * np.median(np.abs(current_window - baseline_window))
        )

        acceptable_fraction = (
            event_fraction <= cfg.representative_trace_max_event_fraction
        )
        if n_events >= cfg.representative_trace_min_events and acceptable_fraction:
            tier = 0
        elif n_events >= 1 and acceptable_fraction:
            tier = 1
        elif acceptable_fraction:
            tier = 2
        else:
            tier = 3

        # This objective ranking avoids choosing a trace by appearance:
        # 1) event-content tier; 2) baseline drift; 3) residual noise;
        # 4) more filtered events; 5) earliest deterministic location.
        rank = (
            tier,
            baseline_drift,
            noise_mad,
            -n_events,
            experiment_index,
            sweep,
            absolute_start,
        )

        candidate = {
            "selection_rank": rank,
            "selection_tier": tier,
            "selection_reason": TRACE_TIER_LABELS[tier],
            "analyte": analyte,
            "pore_id": pore_id,
            "source_pore_folder": source_pore_folder,
            "experiment_index": experiment_index,
            "file_name": file_name,
            "sweep": sweep,
            "fs_Hz": fs_hz,
            "start_idx_in_sweep": int(absolute_start),
            "end_idx_in_sweep": int(absolute_end - 1),
            "start_s_in_sweep": float(absolute_start / fs_hz),
            "end_s_in_sweep": float((absolute_end - 1) / fs_hz),
            "duration_s": float(window_points / fs_hz),
            "n_filtered_events": n_events,
            "event_fraction": event_fraction,
            "baseline_drift_p95_p05_pA": baseline_drift,
            "residual_noise_MAD_sigma_pA": noise_mad,
            "current_pA": current_window.astype(np.float32, copy=True),
            "baseline_pA": baseline_window.astype(np.float32, copy=True),
            "sigma_pA": sigma_window.astype(np.float32, copy=True),
            "threshold_enter_pA": threshold_window.astype(np.float32, copy=True),
            "selected_events": selected_events[[
                "start_idx_in_sweep",
                "end_idx_in_sweep",
            ]].copy(),
        }

        if candidate_is_better(candidate, current_best):
            current_best = candidate

    return current_best


def trace_candidate_to_dataframe(
    candidate: Dict[str, object],
    cfg: DetectConfig,
) -> pd.DataFrame:
    current = np.asarray(candidate["current_pA"], dtype=np.float32)
    baseline = np.asarray(candidate["baseline_pA"], dtype=np.float32)
    sigma = np.asarray(candidate["sigma_pA"], dtype=np.float32)
    threshold_enter = np.asarray(
        candidate["threshold_enter_pA"],
        dtype=np.float32,
    )
    fs_hz = float(candidate["fs_Hz"])
    start_idx = int(candidate["start_idx_in_sweep"])

    event_detected = np.zeros(len(current), dtype=np.int8)
    selected_events = candidate["selected_events"]
    for row in selected_events.itertuples(index=False):
        low = max(0, int(row.start_idx_in_sweep) - start_idx)
        high = min(len(current), int(row.end_idx_in_sweep) - start_idx + 1)
        if high > low:
            event_detected[low:high] = 1

    return pd.DataFrame({
        "time_s": np.arange(len(current), dtype=float) / fs_hz,
        "time_in_sweep_s": (
            np.arange(len(current), dtype=float) + start_idx
        ) / fs_hz,
        "current_pA": current,
        "baseline_pA": baseline,
        "sigma_pA": sigma,
        "threshold_enter_pA": threshold_enter,
        "threshold_exit_pA": baseline - cfg.sigma_exit * sigma,
        "event_detected": event_detected,
    })


def trace_candidate_metadata(
    candidate: Dict[str, object],
    trace_sheet: str,
) -> Dict[str, object]:
    excluded = {
        "selection_rank",
        "current_pA",
        "baseline_pA",
        "sigma_pA",
        "threshold_enter_pA",
        "selected_events",
    }
    row = {
        key: value
        for key, value in candidate.items()
        if key not in excluded
    }
    row["trace_sheet"] = trace_sheet
    row["selection_rule"] = (
        "tier -> lowest baseline p95-p05 drift -> lowest residual MAD noise "
        "-> most filtered events -> earliest window"
    )
    return row


# ---------------------------------------------------------------------------
# FILE AND PEPTIDE PROCESSING
# ---------------------------------------------------------------------------

def summarize_sweep(
    values: np.ndarray,
    baseline_samples: List[np.ndarray],
    sigma_samples: List[np.ndarray],
    events: pd.DataFrame,
    fs_hz: float,
    file_name: str,
    analyte: str,
    pore_id: str,
    source_pore_folder: str,
    experiment_index: int,
    channel: int,
    sweep: int,
    i0_global: float,
    sigma0_global: float,
    window_start_idx: int,
    window_end_idx: int,
    flags: Dict[str, bool],
) -> Dict[str, object]:
    baseline_sample = (
        np.concatenate(baseline_samples)
        if baseline_samples
        else np.array([], dtype=float)
    )
    sigma_sample = (
        np.concatenate(sigma_samples)
        if sigma_samples
        else np.array([], dtype=float)
    )
    duration_s = len(values) / fs_hz if fs_hz > 0 else np.nan
    total_event_time_ms = (
        float(events["dwell_total_ms"].sum())
        if not events.empty
        else 0.0
    )
    event_fraction = (
        total_event_time_ms / (duration_s * 1000.0)
        if duration_s > 0
        else np.nan
    )

    return {
        "experiment_index": experiment_index,
        "file_name": file_name,
        "source_pore_folder": source_pore_folder,
        "pore_id": pore_id,
        "analyte": analyte,
        "channel": channel,
        "sweep": sweep,
        "window_start_idx": int(window_start_idx),
        "window_end_idx": int(window_end_idx),
        "window_start_s": float(window_start_idx / fs_hz),
        "window_end_s": float(window_end_idx / fs_hz),
        "n_samples": int(len(values)),
        "duration_s": float(duration_s),
        "fs_Hz": fs_hz,
        "I0_global_pA": i0_global,
        "sigma0_global_pA": sigma0_global,
        "trace_std_pA": safe_std(values),
        "dynamic_range_pA": (
            float(np.percentile(values, 99) - np.percentile(values, 1))
            if len(values)
            else np.nan
        ),
        "baseline_median_pA": safe_median(baseline_sample),
        "baseline_range_pA": (
            float(np.max(baseline_sample) - np.min(baseline_sample))
            if baseline_sample.size
            else np.nan
        ),
        "sigma_median_pA": safe_median(sigma_sample),
        "n_events": int(len(events)),
        "event_rate_Hz": (
            len(events) / duration_s
            if duration_s > 0
            else np.nan
        ),
        "event_fraction": event_fraction,
        "total_event_time_ms": total_event_time_ms,
        "median_dwell_total_ms": (
            float(events["dwell_total_ms"].median())
            if not events.empty
            else np.nan
        ),
        "median_dwell_core_ms": (
            float(events["dwell_core_ms"].median())
            if not events.empty
            else np.nan
        ),
        "median_blockade_depth_mean_pA": (
            float(events["blockade_depth_mean_pA"].median())
            if not events.empty
            else np.nan
        ),
        "median_blockade_depth_min_pA": (
            float(events["blockade_depth_min_pA"].median())
            if not events.empty
            else np.nan
        ),
        "median_I_over_I0_pct": (
            float(events["I_over_I0_pct"].median())
            if not events.empty
            else np.nan
        ),
        **flags,
    }


def process_abf_file(
    experiment: Dict[str, object],
    cfg: DetectConfig,
) -> Dict[str, object]:
    abf_path = str(experiment["abf_path"])
    experiment_index = int(experiment["experiment_index"])
    analyte = str(experiment["analyte"])
    pore_id = str(experiment["pore_id"])
    source_pore_folder = str(experiment["source_pore_folder"])
    channel = int(experiment["channel"])
    file_name = Path(abf_path).name

    print(f"Reading: {abf_path}")
    print(
        f"  peptide={analyte} | pore={pore_id} | "
        f"source folder={source_pore_folder}"
    )

    abf = pyabf.ABF(abf_path)
    fs_hz = float(abf.dataRate)
    i0_global, sigma0_global, unit_scale = sample_abf_for_i0(
        abf,
        channel,
    )
    print(
        f"  sweeps={abf.sweepCount} | channels={abf.channelCount} | "
        f"fs={fs_hz:.2f} Hz | I0={i0_global:.3f} pA"
    )

    file_event_tables: List[pd.DataFrame] = []
    sweep_rows: List[Dict[str, object]] = []
    best_trace: Optional[Dict[str, object]] = None
    total_analyzed_samples = 0
    global_offset = 0

    for sweep in range(abf.sweepCount):
        abf.setSweep(sweep, channel=channel)
        full_values = (
            np.asarray(abf.sweepY, dtype=np.float32)
            * np.float32(unit_scale)
        )
        if not np.isfinite(full_values).all():
            raise ValueError(
                f"NaN or Inf found in sweep {sweep} of {abf_path}"
            )

        values, window_start_idx, window_end_idx = restrict_sweep_time_window(
            full_values,
            fs_hz,
            cfg,
        )
        total_analyzed_samples += len(values)
        print(
            f"  sweep {sweep + 1}/{abf.sweepCount}: "
            f"{window_start_idx / fs_hz:.3f}-"
            f"{window_end_idx / fs_hz:.3f} s"
        )

        if not len(values):
            flags = {
                "flag_flat_trace": False,
                "flag_large_baseline_drift": False,
                "flag_high_event_fraction": False,
                "flag_no_events": True,
            }
            sweep_rows.append(summarize_sweep(
                values=values,
                baseline_samples=[],
                sigma_samples=[],
                events=empty_events_df(),
                fs_hz=fs_hz,
                file_name=file_name,
                analyte=analyte,
                pore_id=pore_id,
                source_pore_folder=source_pore_folder,
                experiment_index=experiment_index,
                channel=channel,
                sweep=sweep,
                i0_global=i0_global,
                sigma0_global=sigma0_global,
                window_start_idx=window_start_idx,
                window_end_idx=window_end_idx,
                flags=flags,
            ))
            global_offset += len(full_values)
            continue

        chunk_points = max(
            1,
            int(round(DETECTION_CHUNK_DURATION_S * fs_hz)),
        )
        sweep_event_tables: List[pd.DataFrame] = []
        chunk_flags: List[Dict[str, bool]] = []
        baseline_samples: List[np.ndarray] = []
        sigma_samples: List[np.ndarray] = []

        for chunk_start in range(0, len(values), chunk_points):
            chunk_end = min(len(values), chunk_start + chunk_points)
            chunk_values = values[chunk_start:chunk_end]
            absolute_chunk_start = window_start_idx + chunk_start

            (
                chunk_events,
                baseline,
                sigma,
                threshold_enter,
                flags,
            ) = detect_chunk(
                values=chunk_values,
                fs_hz=fs_hz,
                cfg=cfg,
                file_name=file_name,
                analyte=analyte,
                channel=channel,
                sweep=sweep,
                global_offset=global_offset,
                sweep_window_start_idx=absolute_chunk_start,
                i0_global=i0_global,
            )

            if not chunk_events.empty:
                chunk_events = chunk_events.copy()
                chunk_events["pore_id"] = pore_id
                chunk_events["experiment_index"] = experiment_index
                chunk_events["source_pore_folder"] = source_pore_folder
                sweep_event_tables.append(chunk_events)

            best_trace = update_best_trace_from_chunk(
                current_best=best_trace,
                values=chunk_values,
                baseline=baseline,
                sigma=sigma,
                threshold_enter=threshold_enter,
                events=chunk_events,
                fs_hz=fs_hz,
                cfg=cfg,
                analyte=analyte,
                pore_id=pore_id,
                source_pore_folder=source_pore_folder,
                experiment_index=experiment_index,
                file_name=file_name,
                sweep=sweep,
                chunk_start_idx_in_sweep=absolute_chunk_start,
            )

            if len(baseline):
                sample_count = min(1000, len(baseline))
                sample_indices = np.linspace(
                    0,
                    len(baseline) - 1,
                    sample_count,
                    dtype=np.int64,
                )
                baseline_samples.append(baseline[sample_indices])
                sigma_samples.append(sigma[sample_indices])

            chunk_flags.append(flags)
            del baseline, sigma, threshold_enter, chunk_values

        sweep_events = (
            pd.concat(sweep_event_tables, ignore_index=True)
            if sweep_event_tables
            else empty_events_df()
        )
        if not sweep_events.empty:
            file_event_tables.append(sweep_events)

        combined_flags = {
            key: any(flag.get(key, False) for flag in chunk_flags)
            for key in (
                "flag_flat_trace",
                "flag_large_baseline_drift",
                "flag_high_event_fraction",
                "flag_no_events",
            )
        }
        combined_flags["flag_no_events"] = sweep_events.empty

        sweep_rows.append(summarize_sweep(
            values=values,
            baseline_samples=baseline_samples,
            sigma_samples=sigma_samples,
            events=sweep_events,
            fs_hz=fs_hz,
            file_name=file_name,
            analyte=analyte,
            pore_id=pore_id,
            source_pore_folder=source_pore_folder,
            experiment_index=experiment_index,
            channel=channel,
            sweep=sweep,
            i0_global=i0_global,
            sigma0_global=sigma0_global,
            window_start_idx=window_start_idx,
            window_end_idx=window_end_idx,
            flags=combined_flags,
        ))

        global_offset += len(full_values)
        del full_values, values, sweep_event_tables
        gc.collect()

    events = (
        pd.concat(file_event_tables, ignore_index=True)
        if file_event_tables
        else empty_events_df()
    )
    sweep_summary = pd.DataFrame(sweep_rows)
    duration_s = total_analyzed_samples / fs_hz if fs_hz > 0 else np.nan
    total_event_time_ms = (
        float(events["dwell_total_ms"].sum())
        if not events.empty
        else 0.0
    )

    file_summary = {
        "experiment_index": experiment_index,
        "file_name": file_name,
        "source_pore_folder": source_pore_folder,
        "pore_id": pore_id,
        "analyte": analyte,
        "channel": channel,
        "fs_Hz": fs_hz,
        "n_sweeps": int(abf.sweepCount),
        "total_analyzed_samples": int(total_analyzed_samples),
        "total_analyzed_duration_s": float(duration_s),
        "I0_global_pA": i0_global,
        "sigma0_global_pA": sigma0_global,
        "total_events_before_filter": int(len(events)),
        "event_rate_Hz_before_filter": (
            len(events) / duration_s
            if duration_s > 0
            else np.nan
        ),
        "event_fraction_before_filter": (
            total_event_time_ms / (duration_s * 1000.0)
            if duration_s > 0
            else np.nan
        ),
        "median_dwell_total_ms_before_filter": (
            float(events["dwell_total_ms"].median())
            if not events.empty
            else np.nan
        ),
        "median_blockade_depth_mean_pA_before_filter": (
            float(events["blockade_depth_mean_pA"].median())
            if not events.empty
            else np.nan
        ),
        "n_sweeps_with_events": (
            int((sweep_summary["n_events"] > 0).sum())
            if not sweep_summary.empty
            else 0
        ),
        "n_sweeps_flagged_flat": (
            int(sweep_summary["flag_flat_trace"].sum())
            if not sweep_summary.empty
            else 0
        ),
        "n_sweeps_flagged_drift": (
            int(sweep_summary["flag_large_baseline_drift"].sum())
            if not sweep_summary.empty
            else 0
        ),
        "n_sweeps_flagged_high_event_fraction": (
            int(sweep_summary["flag_high_event_fraction"].sum())
            if not sweep_summary.empty
            else 0
        ),
        "n_sweeps_flagged_no_events": (
            int(sweep_summary["flag_no_events"].sum())
            if not sweep_summary.empty
            else 0
        ),
    }

    return {
        "events": events,
        "sweep_summary": sweep_summary,
        "file_summary": file_summary,
        "best_trace": best_trace,
    }


def robust_mad_z(
    values: pd.Series,
    mad_floor: float = 1e-12,
) -> pd.Series:
    numeric = pd.to_numeric(values, errors="coerce")
    median = numeric.median(skipna=True)
    mad = (
        median_abs_deviation(numeric.dropna(), scale="normal")
        if numeric.notna().sum()
        else np.nan
    )
    if not np.isfinite(mad) or mad < mad_floor:
        return pd.Series(np.zeros(len(numeric)), index=values.index, dtype=float)
    return (numeric - median) / mad


def add_between_sweep_qc_flags(
    sweep_summary: pd.DataFrame,
    cfg: DetectConfig,
) -> pd.DataFrame:
    if sweep_summary.empty or not cfg.enable_between_sweep_qc:
        return sweep_summary

    output = sweep_summary.copy()
    group_columns = [
        column
        for column in cfg.qc_group_columns
        if column in output.columns
    ]
    if not group_columns:
        return output

    for metric in cfg.qc_metrics:
        if metric not in output.columns:
            continue
        z_column = f"qc_z_{metric}"
        flag_column = f"flag_sweep_{metric}_outlier"
        output[z_column] = np.nan
        output[flag_column] = False

        for _, indices in output.groupby(
            group_columns,
            dropna=False,
        ).groups.items():
            indices = list(indices)
            if len(indices) < cfg.qc_min_sweeps_per_group:
                continue
            z_scores = robust_mad_z(
                output.loc[indices, metric],
                cfg.qc_min_mad_floor,
            )
            output.loc[indices, z_column] = z_scores
            output.loc[indices, flag_column] = (
                z_scores.abs() > cfg.qc_mad_z_threshold
            )

    flag_columns = [
        column
        for column in output.columns
        if column.startswith("flag_sweep_")
        and column.endswith("_outlier")
    ]
    output["flag_sweep_any_between_sweep_outlier"] = (
        output[flag_columns].any(axis=1)
        if flag_columns
        else False
    )
    return output


def update_file_summary_with_filtered_events(
    file_summary: pd.DataFrame,
    filtered_events: pd.DataFrame,
) -> pd.DataFrame:
    output = file_summary.copy()
    if output.empty:
        return output

    metric_columns = [
        "n_events_after_filter",
        "median_dwell_total_ms_after_filter",
        "median_blockade_depth_mean_pA_after_filter",
        "median_I_over_I0_pct_after_filter",
    ]
    if filtered_events.empty:
        output["n_events_after_filter"] = 0
        for column in metric_columns[1:]:
            output[column] = np.nan
    else:
        metrics = (
            filtered_events.groupby("experiment_index", dropna=False)
            .agg(
                n_events_after_filter=("event_id", "count"),
                median_dwell_total_ms_after_filter=(
                    "dwell_total_ms",
                    "median",
                ),
                median_blockade_depth_mean_pA_after_filter=(
                    "blockade_depth_mean_pA",
                    "median",
                ),
                median_I_over_I0_pct_after_filter=(
                    "I_over_I0_pct",
                    "median",
                ),
            )
            .reset_index()
        )
        output = output.merge(metrics, on="experiment_index", how="left")

    output["n_events_after_filter"] = (
        pd.to_numeric(
            output["n_events_after_filter"],
            errors="coerce",
        )
        .fillna(0)
        .astype(int)
    )
    output["event_rate_Hz_after_filter"] = (
        output["n_events_after_filter"]
        / output["total_analyzed_duration_s"].replace(0, np.nan)
    )
    return output


def build_pore_summary(
    analyte: str,
    pore_id: str,
    source_pore_folder: str,
    experiments: Sequence[Dict[str, object]],
    events_unfiltered: pd.DataFrame,
    events_filtered: pd.DataFrame,
    file_summary: pd.DataFrame,
    trace_sheet: str,
) -> Dict[str, object]:
    duration_s = (
        float(file_summary["total_analyzed_duration_s"].sum())
        if not file_summary.empty
        else 0.0
    )
    return {
        "analyte": analyte,
        "pore_id": pore_id,
        "source_pore_folder": source_pore_folder,
        "n_direct_abf_files": len(experiments),
        "n_successfully_processed_abf_files": int(len(file_summary)),
        "n_sweeps": (
            int(file_summary["n_sweeps"].sum())
            if not file_summary.empty
            else 0
        ),
        "analyzed_duration_s": duration_s,
        "analyzed_duration_min": duration_s / 60.0,
        "n_events_before_filter": int(len(events_unfiltered)),
        "n_events_after_filter": int(len(events_filtered)),
        "event_rate_Hz_after_filter": (
            len(events_filtered) / duration_s
            if duration_s > 0
            else np.nan
        ),
        "median_dwell_total_ms_after_filter": (
            float(events_filtered["dwell_total_ms"].median())
            if not events_filtered.empty
            else np.nan
        ),
        "median_blockade_depth_mean_pA_after_filter": (
            float(events_filtered["blockade_depth_mean_pA"].median())
            if not events_filtered.empty
            else np.nan
        ),
        "median_I_over_I0_pct_after_filter": (
            float(events_filtered["I_over_I0_pct"].median())
            if not events_filtered.empty
            else np.nan
        ),
        "median_I0_global_pA": (
            float(file_summary["I0_global_pA"].median())
            if not file_summary.empty
            else np.nan
        ),
        "median_sigma0_global_pA": (
            float(file_summary["sigma0_global_pA"].median())
            if not file_summary.empty
            else np.nan
        ),
        "trace_sheet": trace_sheet,
    }


# ---------------------------------------------------------------------------
# MEMORY-SAFE EXCEL WRITING
# ---------------------------------------------------------------------------

def excel_value(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not np.isfinite(value):
        return None
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if pd.isna(value):
        return None
    return value


def add_bold_header(
    worksheet,
    columns: Sequence[object],
) -> None:
    cells = []
    for column in columns:
        cell = WriteOnlyCell(worksheet, value=str(column))
        cell.font = Font(bold=True)
        cells.append(cell)
    worksheet.append(cells)

    for index, column in enumerate(columns, start=1):
        width = min(32, max(11, len(str(column)) + 2))
        worksheet.column_dimensions[get_column_letter(index)].width = width


def append_dataframe_sheet(
    workbook: Workbook,
    sheet_name: str,
    dataframe: pd.DataFrame,
) -> str:
    sheet_name = sheet_name[:31]
    worksheet = workbook.create_sheet(sheet_name)
    columns = list(dataframe.columns)
    add_bold_header(worksheet, columns)
    for row in dataframe.itertuples(index=False, name=None):
        worksheet.append([excel_value(value) for value in row])
    return sheet_name


def iter_pickled_frames(paths: Sequence[Path]) -> Iterable[pd.DataFrame]:
    for path in paths:
        frame = pd.read_pickle(path)
        yield frame
        del frame
        gc.collect()


def append_pickled_event_sheets(
    workbook: Workbook,
    base_sheet_name: str,
    paths: Sequence[Path],
    columns: Sequence[str],
) -> List[str]:
    sheet_names: List[str] = []
    worksheet = workbook.create_sheet(base_sheet_name[:31])
    sheet_names.append(worksheet.title)
    add_bold_header(worksheet, columns)
    rows_in_sheet = 0
    part_number = 1

    for dataframe in iter_pickled_frames(paths):
        frame = dataframe.reindex(columns=columns)
        for row in frame.itertuples(index=False, name=None):
            if rows_in_sheet >= EXCEL_MAX_DATA_ROWS:
                part_number += 1
                next_name = f"{base_sheet_name}_{part_number:02d}"[:31]
                worksheet = workbook.create_sheet(next_name)
                sheet_names.append(worksheet.title)
                add_bold_header(worksheet, columns)
                rows_in_sheet = 0

            worksheet.append([excel_value(value) for value in row])
            rows_in_sheet += 1

        del dataframe, frame
        gc.collect()

    return sheet_names


def parameter_table(cfg: DetectConfig) -> pd.DataFrame:
    rows: List[Dict[str, object]] = [
        {
            "scope": "run",
            "parameter": "DATA_ROOT",
            "value": DATA_ROOT,
        },
        {
            "scope": "run",
            "parameter": "OUTPUT_ROOT_NAME",
            "value": OUTPUT_ROOT_NAME,
        },
        {
            "scope": "run",
            "parameter": "CHANNEL",
            "value": CHANNEL,
        },
        {
            "scope": "run",
            "parameter": "DETECTION_CHUNK_DURATION_S",
            "value": DETECTION_CHUNK_DURATION_S,
        },
        {
            "scope": "discovery",
            "parameter": "folder_rule",
            "value": (
                "one immediate DATA_ROOT child folder = one pore; "
                "only direct ABFs; never recursive"
            ),
        },
        {
            "scope": "trace",
            "parameter": "signal_export",
            "value": (
                "same ABF samples analysed by detection; converted nA to pA "
                "only when unit heuristic triggers; no added digital filter"
            ),
        },
    ]
    for key, value in asdict(cfg).items():
        rows.append({
            "scope": "DetectConfig",
            "parameter": key,
            "value": value,
        })
    return pd.DataFrame(rows)


def readme_table(
    analyte: str,
    filtered_sheet_names: Sequence[str],
    unfiltered_sheet_names: Sequence[str],
) -> pd.DataFrame:
    return pd.DataFrame([
        {
            "item": "peptide",
            "details": analyte,
        },
        {
            "item": "workbook scope",
            "details": (
                "This workbook contains one physical pore only. Multiple ABFs "
                "are combined only when they are directly inside the same "
                "source pore folder."
            ),
        },
        {
            "item": "input discovery",
            "details": (
                "Each immediate folder in Recordings is one physical pore. "
                "Only ABFs directly inside it are used; IV Curve, other "
                "channels, baseline and every other subfolder are ignored."
            ),
        },
        {
            "item": "events",
            "details": (
                "Filtered event table used for R. Sheet(s): "
                + ", ".join(filtered_sheet_names)
            ),
        },
        {
            "item": "events_unfiltered",
            "details": (
                "All events passing the detector's 0.05-500 ms limits, before "
                "the output filters. Sheet(s): "
                + ", ".join(unfiltered_sheet_names)
            ),
        },
        {
            "item": "current filters in events",
            "details": (
                f"dwell {CFG.min_dwell_filter_ms}-{CFG.max_dwell_filter_ms} ms; "
                f"{CFG.depth_filter_column} "
                f"{CFG.min_depth_filter_pA}-"
                f"{CFG.max_depth_filter_pA} pA; "
                f"I/I0 filter enabled={CFG.filter_by_i_over_i0_pct}"
            ),
        },
        {
            "item": "trace_Pn",
            "details": (
                "One representative raw-current segment per physical pore, "
                "with the baseline, sigma, thresholds and event mask used by "
                "the detector. Details are in trace_selection."
            ),
        },
        {
            "item": "trace selection",
            "details": (
                "All candidate windows are ranked deterministically: event "
                "tier, lowest baseline p95-p05 drift, lowest residual MAD "
                "noise, most filtered events, then earliest window."
            ),
        },
        {
            "item": "R import",
            "details": (
                "Use readxl::excel_sheets(path) and readxl::read_excel(path, "
                "sheet = ...). Read all sheets matching ^events($|_[0-9]+$) "
                "if the events table was split at Excel's row limit."
            ),
        },
    ])


def save_small_excel(
    path: Path,
    sheets: Sequence[Tuple[str, pd.DataFrame]],
) -> None:
    workbook = Workbook(write_only=True)
    for sheet_name, dataframe in sheets:
        append_dataframe_sheet(workbook, sheet_name, dataframe)
    temporary_path = path.with_name(path.stem + "_writing.xlsx")
    workbook.save(temporary_path)
    os.replace(temporary_path, path)


def process_one_pore(
    analyte: str,
    experiments: Sequence[Dict[str, object]],
    analyte_audit: pd.DataFrame,
    output_dir: Path,
    cfg: DetectConfig,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    experiments_by_pore: Dict[str, List[Dict[str, object]]] = {}
    for experiment in experiments:
        experiments_by_pore.setdefault(
            str(experiment["pore_id"]),
            [],
        ).append(experiment)

    pore_ids = sorted(experiments_by_pore, key=natural_sort_key)
    if len(pore_ids) != 1:
        raise ValueError(
            "process_one_pore requires exactly one physical pore, "
            f"but received: {pore_ids}"
        )

    workbook_pore_id = pore_ids[0]
    output_path = output_dir / (
        f"{safe_name(analyte)}_{safe_name(workbook_pore_id)}"
        "_nanopore_R_data.xlsx"
    )

    sweep_frames: List[pd.DataFrame] = []
    file_frames: List[pd.DataFrame] = []
    pore_rows: List[Dict[str, object]] = []
    trace_candidates: Dict[str, Dict[str, object]] = {}
    error_rows: List[Dict[str, object]] = []
    filtered_paths: List[Path] = []
    unfiltered_paths: List[Path] = []
    next_raw_event_id = 1
    next_filtered_event_id = 1

    with tempfile.TemporaryDirectory(
        prefix=f"nanopore_{safe_name(analyte)}_",
    ) as temporary_directory:
        temporary_root = Path(temporary_directory)

        for pore_id in pore_ids:
            pore_experiments = sorted(
                experiments_by_pore[pore_id],
                key=lambda item: natural_sort_key(item["abf_path"]),
            )
            source_pore_folder = str(
                pore_experiments[0]["source_pore_folder"]
            )
            print()
            print(
                f"{analyte} | {pore_id} | {source_pore_folder} | "
                f"{len(pore_experiments)} direct ABF file(s)"
            )

            pore_event_frames: List[pd.DataFrame] = []
            pore_sweep_frames: List[pd.DataFrame] = []
            pore_file_rows: List[Dict[str, object]] = []
            pore_best_trace: Optional[Dict[str, object]] = None

            for experiment in pore_experiments:
                try:
                    result = process_abf_file(experiment, cfg)
                except Exception as exc:
                    print(f"  ERROR: {exc}")
                    error_rows.append({
                        "analyte": analyte,
                        "pore_id": pore_id,
                        "source_pore_folder": source_pore_folder,
                        "abf_file": Path(str(experiment["abf_path"])).name,
                        "abf_path": str(experiment["abf_path"]),
                        "error_type": type(exc).__name__,
                        "error_message": str(exc),
                        "traceback": traceback.format_exc(),
                    })
                    continue

                if not result["events"].empty:
                    pore_event_frames.append(result["events"])
                if not result["sweep_summary"].empty:
                    pore_sweep_frames.append(result["sweep_summary"])
                pore_file_rows.append(result["file_summary"])

                file_best = result["best_trace"]
                if (
                    file_best is not None
                    and candidate_is_better(file_best, pore_best_trace)
                ):
                    pore_best_trace = file_best

                del result
                gc.collect()

            pore_events_unfiltered = (
                pd.concat(pore_event_frames, ignore_index=True)
                if pore_event_frames
                else pd.DataFrame(columns=FINAL_EVENT_COLUMNS)
            )
            if not pore_events_unfiltered.empty:
                count = len(pore_events_unfiltered)
                pore_events_unfiltered["raw_event_id"] = np.arange(
                    next_raw_event_id,
                    next_raw_event_id + count,
                )
                next_raw_event_id += count
                pore_events_unfiltered["event_id"] = (
                    pore_events_unfiltered["raw_event_id"]
                )

            pore_events_unfiltered = pore_events_unfiltered.reindex(
                columns=FINAL_EVENT_COLUMNS,
            )
            pore_events_filtered = filter_events_for_outputs(
                pore_events_unfiltered,
                cfg,
            )
            if not pore_events_filtered.empty:
                count = len(pore_events_filtered)
                pore_events_filtered["event_id"] = np.arange(
                    next_filtered_event_id,
                    next_filtered_event_id + count,
                )
                next_filtered_event_id += count
            pore_events_filtered = pore_events_filtered.reindex(
                columns=FINAL_EVENT_COLUMNS,
            )

            pore_file_summary = pd.DataFrame(pore_file_rows)
            pore_file_summary = update_file_summary_with_filtered_events(
                pore_file_summary,
                pore_events_filtered,
            )
            if not pore_file_summary.empty:
                file_frames.append(pore_file_summary)
            if pore_sweep_frames:
                sweep_frames.extend(pore_sweep_frames)

            trace_sheet = ""
            if pore_best_trace is not None:
                trace_sheet = f"trace_{pore_id}"[:31]
                trace_candidates[pore_id] = pore_best_trace

            pore_rows.append(build_pore_summary(
                analyte=analyte,
                pore_id=pore_id,
                source_pore_folder=source_pore_folder,
                experiments=pore_experiments,
                events_unfiltered=pore_events_unfiltered,
                events_filtered=pore_events_filtered,
                file_summary=pore_file_summary,
                trace_sheet=trace_sheet,
            ))

            unfiltered_path = temporary_root / f"{pore_id}_unfiltered.pkl"
            filtered_path = temporary_root / f"{pore_id}_filtered.pkl"
            pore_events_unfiltered.to_pickle(unfiltered_path)
            pore_events_filtered.to_pickle(filtered_path)
            unfiltered_paths.append(unfiltered_path)
            filtered_paths.append(filtered_path)

            del (
                pore_event_frames,
                pore_events_unfiltered,
                pore_events_filtered,
                pore_file_summary,
            )
            gc.collect()

        sweep_summary = (
            pd.concat(sweep_frames, ignore_index=True)
            if sweep_frames
            else pd.DataFrame()
        )
        sweep_summary = add_between_sweep_qc_flags(sweep_summary, cfg)
        file_summary = (
            pd.concat(file_frames, ignore_index=True)
            if file_frames
            else pd.DataFrame()
        )
        pore_summary = pd.DataFrame(pore_rows)
        processing_errors = pd.DataFrame(error_rows, columns=[
            "analyte",
            "pore_id",
            "source_pore_folder",
            "abf_file",
            "abf_path",
            "error_type",
            "error_message",
            "traceback",
        ])
        parameters = parameter_table(cfg)

        trace_metadata_rows: List[Dict[str, object]] = []
        trace_frames: List[Tuple[str, pd.DataFrame]] = []
        for pore_id in sorted(trace_candidates, key=natural_sort_key):
            candidate = trace_candidates[pore_id]
            trace_sheet = f"trace_{pore_id}"[:31]
            trace_metadata_rows.append(
                trace_candidate_metadata(candidate, trace_sheet)
            )
            trace_frames.append((
                trace_sheet,
                trace_candidate_to_dataframe(candidate, cfg),
            ))
        trace_selection = pd.DataFrame(trace_metadata_rows)

        workbook = Workbook(write_only=True)

        # Event sheets are written first in a streaming manner. The README is
        # added afterwards because it records the actual split-sheet names.
        filtered_sheet_names = append_pickled_event_sheets(
            workbook,
            "events",
            filtered_paths,
            FINAL_EVENT_COLUMNS,
        )
        unfiltered_sheet_names = append_pickled_event_sheets(
            workbook,
            "events_unfiltered",
            unfiltered_paths,
            FINAL_EVENT_COLUMNS,
        )

        append_dataframe_sheet(
            workbook,
            "README",
            readme_table(
                analyte,
                filtered_sheet_names,
                unfiltered_sheet_names,
            ),
        )
        append_dataframe_sheet(workbook, "pore_summary", pore_summary)
        append_dataframe_sheet(workbook, "sweep_summary", sweep_summary)
        append_dataframe_sheet(workbook, "file_summary", file_summary)
        append_dataframe_sheet(
            workbook,
            "trace_selection",
            trace_selection,
        )
        append_dataframe_sheet(workbook, "parameters", parameters)
        append_dataframe_sheet(
            workbook,
            "discovery_audit",
            analyte_audit,
        )
        append_dataframe_sheet(
            workbook,
            "processing_errors",
            processing_errors,
        )

        for trace_sheet, trace_frame in trace_frames:
            append_dataframe_sheet(workbook, trace_sheet, trace_frame)

        temporary_output = output_path.with_name(
            output_path.stem + "_writing.xlsx"
        )
        workbook.save(temporary_output)
        os.replace(temporary_output, output_path)

    print(f"Saved: {output_path}")
    return output_path


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main() -> None:
    data_root = Path(DATA_ROOT)
    if not data_root.is_dir():
        raise FileNotFoundError(
            f"DATA_ROOT not found:\n{DATA_ROOT}\n\n"
            "Edit DATA_ROOT near the top of the script."
        )

    print("=" * 80)
    print("NANOPORE DETECTION v33 - ONE EXCEL WORKBOOK PER PORE")
    print(f"DATA_ROOT: {data_root}")
    print("Discovery is non-recursive: only direct ABFs are used.")
    print("=" * 80)

    experiments_by_analyte, discovery_audit = discover_recordings(data_root)
    output_root = data_root / OUTPUT_ROOT_NAME
    output_root.mkdir(parents=True, exist_ok=True)

    audit_path = output_root / "00_ABF_discovery_audit.xlsx"
    save_small_excel(
        audit_path,
        [("discovery_audit", discovery_audit)],
    )
    print(f"Saved discovery audit: {audit_path}")

    if not experiments_by_analyte:
        print(
            "No usable direct ABF files were found. Check the discovery audit "
            "and PEPTIDE_ALIASES."
        )
        return

    total_pores = 0
    total_abfs = 0
    for analyte in sorted(experiments_by_analyte, key=natural_sort_key):
        experiments = experiments_by_analyte[analyte]
        pore_count = len({str(item["pore_id"]) for item in experiments})
        total_pores += pore_count
        total_abfs += len(experiments)

        print()
        print("#" * 80)
        print(
            f"PEPTIDE: {analyte} | "
            f"{pore_count} pore(s) | {len(experiments)} direct ABF(s)"
        )
        print("#" * 80)

        analyte_audit = discovery_audit.loc[
            discovery_audit["analyte"] == analyte
        ].reset_index(drop=True)

        experiments_by_pore: Dict[str, List[Dict[str, object]]] = {}
        for experiment in experiments:
            experiments_by_pore.setdefault(
                str(experiment["pore_id"]),
                [],
            ).append(experiment)

        for pore_id in sorted(experiments_by_pore, key=natural_sort_key):
            pore_experiments = experiments_by_pore[pore_id]
            pore_audit = analyte_audit.loc[
                analyte_audit["pore_id"].astype(str) == pore_id
            ].reset_index(drop=True)

            process_one_pore(
                analyte=analyte,
                experiments=pore_experiments,
                analyte_audit=pore_audit,
                output_dir=output_root / safe_name(analyte),
                cfg=CFG,
            )

    print()
    print("=" * 80)
    print(
        f"FINISHED: {len(experiments_by_analyte)} peptide(s), "
        f"{total_pores} physical pore(s), {total_abfs} direct ABF file(s)."
    )
    print(f"Outputs: {output_root}")
    print("=" * 80)


if __name__ == "__main__":
    main()

